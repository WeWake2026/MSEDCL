import io, zipfile, sqlite3
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def apply_excel_styling(ws, title="Report"):
    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Freeze top row
    ws.freeze_panes = "A2"
    
    # Style header row
    for col_num in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    # Style data rows and auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 1:
                cell.font = data_font
                cell.border = thin_border
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)
        
    # Enable autofilter on the table range
    ws.auto_filter.ref = ws.dimensions

def dataframe_to_styled_excel(df, sheet_name="Data"):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
        ws = writer.sheets[sheet_name[:31]]
        apply_excel_styling(ws, sheet_name)
    output.seek(0)
    return output

def generate_combined_excel(sheets_dict):
    """
    sheets_dict: {'Sheet Name': DataFrame, ...}
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for sheet_name, df in sheets_dict.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)
            ws = writer.sheets[safe_name]
            apply_excel_styling(ws, safe_name)
    output.seek(0)
    return output

def build_export_datasets(con, start_date=None, end_date=None, department=None, auditor=None, current_advisor_id=None):
    where = ["1=1"]
    params = []
    
    if department and department in ("Inbound", "Outbound"):
        where.append("department = ?")
        params.append(department)
    if start_date:
        where.append("COALESCE(call_date, evaluation_date) >= ?")
        params.append(start_date)
    if end_date:
        where.append("COALESCE(call_date, evaluation_date) <= ?")
        params.append(end_date)
    if auditor:
        where.append("auditor_name = ?")
        params.append(auditor)
    if current_advisor_id:
        where.append("advisor_id = ?")
        params.append(current_advisor_id)
        
    base_where = " AND ".join(where)
    
    # 1. Inbound Audits
    inbound_df = pd.read_sql_query(f"""
        SELECT call_id, call_date, evaluation_date, advisor_id, advisor_name, team_leader,
               auditor_name, quality_score, is_fatal, fatal_category, error_categories, observation, audio_url, source_file
        FROM audits WHERE department='Inbound' AND {base_where.replace('department = ?', '1=1')}
        ORDER BY COALESCE(call_date, evaluation_date) DESC
    """, con, params=[p for i, p in enumerate(params) if not (department and department in ('Inbound','Outbound') and params[i] == department)])
    
    # 2. Outbound Audits
    outbound_df = pd.read_sql_query(f"""
        SELECT call_id, call_date, evaluation_date, advisor_id, advisor_name, team_leader,
               auditor_name, quality_score, is_fatal, fatal_category, error_categories, observation, audio_url, source_file
        FROM audits WHERE department='Outbound' AND {base_where.replace('department = ?', '1=1')}
        ORDER BY COALESCE(call_date, evaluation_date) DESC
    """, con, params=[p for i, p in enumerate(params) if not (department and department in ('Inbound','Outbound') and params[i] == department)])

    # 3. All Filtered Audits
    all_audits_df = pd.read_sql_query(f"""
        SELECT department, call_id, call_date, evaluation_date, advisor_id, advisor_name, team_leader,
               auditor_name, quality_score, is_fatal, fatal_category, error_categories, observation, audio_url, source_file
        FROM audits WHERE {base_where}
        ORDER BY COALESCE(call_date, evaluation_date) DESC
    """, con, params=params)

    # 4. Fatal Calls
    fatal_df = pd.read_sql_query(f"""
        SELECT department, call_id, call_date, evaluation_date, advisor_id, advisor_name, team_leader,
               auditor_name, quality_score, fatal_category, error_categories, observation, audio_url
        FROM audits WHERE is_fatal=1 AND {base_where}
        ORDER BY COALESCE(call_date, evaluation_date) DESC
    """, con, params=params)

    # 5. Low Score (<80%)
    low_df = pd.read_sql_query(f"""
        SELECT department, call_id, call_date, evaluation_date, advisor_id, advisor_name, team_leader,
               auditor_name, quality_score, is_fatal, fatal_category, observation, audio_url
        FROM audits WHERE quality_score < 80 AND {base_where}
        ORDER BY quality_score ASC
    """, con, params=params)

    # 6. Average Score (80% - 90%)
    avg_df = pd.read_sql_query(f"""
        SELECT department, call_id, call_date, evaluation_date, advisor_id, advisor_name, team_leader,
               auditor_name, quality_score, is_fatal, fatal_category, observation, audio_url
        FROM audits WHERE quality_score >= 80 AND quality_score <= 90 AND {base_where}
        ORDER BY quality_score ASC
    """, con, params=params)

    # 7. Good Calls (>90%)
    good_df = pd.read_sql_query(f"""
        SELECT department, call_id, call_date, evaluation_date, advisor_id, advisor_name, team_leader,
               auditor_name, quality_score, is_fatal, fatal_category, observation, audio_url
        FROM audits WHERE quality_score > 90 AND {base_where}
        ORDER BY quality_score DESC
    """, con, params=params)

    # 8. Auditor Summary
    auditor_df = pd.read_sql_query(f"""
        SELECT auditor_name, COUNT(*) as total_audits,
               ROUND(AVG(quality_score), 2) as average_score,
               SUM(is_fatal) as fatal_calls,
               SUM(CASE WHEN quality_score < 80 THEN 1 ELSE 0 END) as low_scores,
               SUM(CASE WHEN quality_score > 90 THEN 1 ELSE 0 END) as good_scores
        FROM audits WHERE {base_where} AND auditor_name IS NOT NULL
        GROUP BY auditor_name ORDER BY total_audits DESC
    """, con, params=params)

    # 9. Advisor Summary
    advisor_df = pd.read_sql_query(f"""
        SELECT advisor_id, advisor_name, department, COUNT(*) as total_audits,
               ROUND(AVG(quality_score), 2) as avg_score,
               SUM(is_fatal) as fatal_count,
               SUM(CASE WHEN quality_score < 80 THEN 1 ELSE 0 END) as low_score_count,
               SUM(CASE WHEN quality_score >= 80 AND quality_score <= 90 THEN 1 ELSE 0 END) as mid_score_count,
               SUM(CASE WHEN quality_score > 90 THEN 1 ELSE 0 END) as good_score_count
        FROM audits WHERE {base_where}
        GROUP BY advisor_id, advisor_name, department
        ORDER BY fatal_count DESC, avg_score ASC
    """, con, params=params)

    # 10. Warning & PIP History
    warnings_df = pd.read_sql_query("""
        SELECT id, advisor_id, advisor_name, department, warning_type, issue_date, effective_date,
               review_date, recipient_email, subject, email_status, status, created_by, approved_by,
               closure_date, remarks
        FROM warnings ORDER BY id DESC
    """, con)

    # 11. Email History
    email_df = pd.read_sql_query("""
        SELECT id, warning_id, created_at, employee_id, advisor_name, department, recipient_email,
               email_type, subject, sent_status, sent_at, delivery_status, error_details
        FROM email_history ORDER BY id DESC
    """, con)

    # 12. Import History
    imports_df = pd.read_sql_query("""
        SELECT id as batch_id, department, source_file, uploaded_by, uploaded_at, rows_seen,
               rows_inserted, duplicates, rejected
        FROM import_batches ORDER BY uploaded_at DESC
    """, con)

    return {
        "Inbound Audits": inbound_df,
        "Outbound Audits": outbound_df,
        "All Audits": all_audits_df,
        "Fatal Calls": fatal_df,
        "Low Score (<80%)": low_df,
        "Average Score (80-90%)": avg_df,
        "Good Calls (>90%)": good_df,
        "Auditor Summary": auditor_df,
        "Advisor Summary": advisor_df,
        "Warning & PIP History": warnings_df,
        "Email History": email_df,
        "Import History": imports_df
    }

def sanitize_for_export(df):
    """
    Sanitize dataframe string columns to neutralize formula injection (CWE-1236).
    Prepends single quote if string starts with =, +, -, @, \\t, or \\r.
    """
    safe_df = df.copy()
    for col in safe_df.columns:
        if safe_df[col].dtype in (object, "string", "str") or str(safe_df[col].dtype).startswith("str"):
            safe_df[col] = safe_df[col].apply(
                lambda val: f"'{val}" if isinstance(val, str) and val and val[0] in ("=", "+", "-", "@", "\t", "\r") else val
            )
    return safe_df

def create_all_reports_zip(datasets):
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        # Add individual CSVs and Excel files
        for name, df in datasets.items():
            clean_name = name.lower().replace(" ", "_").replace("(", "").replace(")", "").replace("<", "lt").replace(">", "gt").replace("-", "_").replace("%", "pct")
            
            sanitized_df = sanitize_for_export(df)

            # CSV
            csv_bytes = sanitized_df.to_csv(index=False).encode('utf-8')
            zf.writestr(f"csv/{clean_name}.csv", csv_bytes)
            
            # Excel
            excel_bytes = dataframe_to_styled_excel(df, sheet_name=name).getvalue()
            zf.writestr(f"excel/{clean_name}.xlsx", excel_bytes)
            
        # Add Master Combined Workbook
        combined_bytes = generate_combined_excel({
            "Inbound Audits": datasets["Inbound Audits"],
            "Outbound Audits": datasets["Outbound Audits"],
            "Fatal Calls": datasets["Fatal Calls"],
            "Low Score Calls": datasets["Low Score (<80%)"],
            "Good Calls": datasets["Good Calls (>90%)"],
            "Auditor Summary": datasets["Auditor Summary"],
            "Advisor Summary": datasets["Advisor Summary"],
            "Warning and PIP History": datasets["Warning & PIP History"],
            "Email History": datasets["Email History"]
        }).getvalue()
        zf.writestr("MSEDCL_Quality_Audit_Combined_Master.xlsx", combined_bytes)
        
    zip_buffer.seek(0)
    return zip_buffer
